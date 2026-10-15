"""Execution review 2026-10-15 — readers (#5 pptx notes/comments, #6 UTF-16, #11 xlsx formulas) and the automatic
model choice (#8 degraded live lists retire nothing, #9 price before version and no cross-provider version
comparison, #10 no capable model ⇒ cheaper fast model, logged)."""
from __future__ import annotations

import io
import logging
import zipfile

import pytest
from sqlalchemy import select


# ---------------------------------------------------------------------------- #5 pptx: speaker notes + comments
def _pptx():
    buf = io.BytesIO()
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        for i in (1, 2):
            z.writestr(f"ppt/slides/slide{i}.xml", f"<p:sld {ns}><a:p><a:r><a:t>SLIDE_{i}</a:t></a:r></a:p></p:sld>")
        # notes numbering deliberately NOT the slide numbering: slide 1 ⇒ notesSlide7, slide 2 ⇒ notesSlide3
        z.writestr("ppt/slides/_rels/slide1.xml.rels",
                   '<Relationships><Relationship Id="r1" Target="../notesSlides/notesSlide7.xml"/>'
                   '<Relationship Id="r2" Target="../comments/comment1.xml"/></Relationships>')
        z.writestr("ppt/slides/_rels/slide2.xml.rels", '<Relationships><Relationship Id="r1" Target="../notesSlides/notesSlide3.xml"/></Relationships>')
        z.writestr("ppt/notesSlides/notesSlide7.xml", f"<p:notes {ns}><a:p><a:r><a:t>NOTES_FOR_1</a:t></a:r></a:p></p:notes>")
        z.writestr("ppt/notesSlides/notesSlide3.xml", f"<p:notes {ns}><a:p><a:r><a:t>NOTES_FOR_2</a:t></a:r></a:p></p:notes>")
        z.writestr("ppt/comments/comment1.xml", f'<p:cmLst {ns}><p:cm authorId="0"><p:text>LEGACY_COMMENT</p:text></p:cm></p:cmLst>')
        z.writestr("ppt/comments/modernComment_101_ABC.xml",
                   f'<p188:cmLst xmlns:p188="x" {ns}><p188:cm><p188:txBody><a:p><a:r><a:t>MODERN_COMMENT</a:t></a:r></a:p></p188:txBody></p188:cm></p188:cmLst>')
    return buf.getvalue()


def test_pptx_reads_speaker_notes_and_comments():
    from app.services.inspection_files import extract

    ex = extract(_pptx(), "deck.pptx", "")
    t = ex["text"]
    assert ex["status"] == "ok"
    for m in ("SLIDE_1", "SLIDE_2", "NOTES_FOR_1", "NOTES_FOR_2", "LEGACY_COMMENT", "MODERN_COMMENT"):
        assert m in t, m
    # each note sits under ITS slide (through the rels), not by file number
    assert t.index("NOTES_FOR_1") < t.index("SLIDE_2") < t.index("NOTES_FOR_2")
    assert t.index("LEGACY_COMMENT") < t.index("SLIDE_2")


# ---------------------------------------------------------------------------- #6 UTF-16 and garbage
FA = "این یک متنِ فارسی است.\nسطرِ دوم با English و 123.\n" * 20


@pytest.mark.parametrize("data", [
    FA.encode("utf-16"),            # with BOM (LE)
    b"\xfe\xff" + FA.encode("utf-16-be"),
    FA.encode("utf-16-le"),         # BOM-less, NUL-heavy
    FA.encode("utf-16-be"),
    ("سلامدنیااینیکمتنبدونفاصلهاست" * 30).encode("utf-16-le"),  # BOM-less, no spaces at all
])
def test_utf16_text_is_decoded_not_reported_as_cp1256_garbage(data):
    from app.services.inspection_files import extract

    ex = extract(data, "notes.txt", "text/plain")
    assert ex["status"] == "ok" and ex["truncated"] is False
    assert "این یک متنِ فارسی است" in ex["text"] or "سلامدنیا" in ex["text"]
    assert "\x00" not in ex["text"]


def test_undecodable_bytes_are_never_reported_complete():
    from app.services.inspection_files import GARBLED_NOTE, extract

    ex = extract(bytes(range(1, 32)) * 200 + b"abc", "dump.log", "text/plain")
    assert ex["truncated"] is True and GARBLED_NOTE in ex["note"]
    assert extract("a,b\n1,2\n".encode("utf-16"), "t.csv", "text/csv")["text"].startswith("a | b")


def test_plain_cp1256_and_utf8_still_read():
    from app.services.inspection_files import extract

    for data in (FA.encode("utf-8"), FA.encode("cp1256", "ignore")):
        ex = extract(data, "a.txt", "text/plain")
        assert ex["status"] == "ok" and ex["truncated"] is False and "فارس" in ex["text"]


# ---------------------------------------------------------------------------- #11 xlsx formula with no cached value
def test_xlsx_formula_without_cached_value_is_read_as_the_formula():
    import openpyxl

    from app.services.inspection_files import extract

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "حساب"
    ws.append([1])
    ws.append([2])
    ws.append([3])
    ws["A4"] = "=SUM(A1:A3)"
    ws["B1"] = "plain"
    buf = io.BytesIO()
    wb.save(buf)
    t = extract(buf.getvalue(), "calc.xlsx", "")["text"]
    assert "=SUM(A1:A3)" in t and "plain" in t and "1" in t


# ---------------------------------------------------------------------------- #8 degraded live lists
def _g(mid, methods=("generateContent",)):
    return {"id": mid, "display": mid, "methods": list(methods), "input_limit": 1048576, "output_limit": 65536}


async def _gemini(db, monkeypatch):
    from app.models.ai_catalog import AICatalogProvider
    from app.services.ai.catalog import seed_ai_catalog

    await seed_ai_catalog(db)
    monkeypatch.setenv("GEMINI_API_KEY", "g-test-123456")
    (await db.get(AICatalogProvider, "gemini")).enabled = True
    await db.commit()


async def _refresh(db, monkeypatch, ids):
    from app.services.ai import model_discovery as md

    async def fake(provider, key):
        return [_g(i) if not i.startswith("text-embedding") else _g(i, ("embedContent",)) for i in ids]

    monkeypatch.setattr(md, "list_live_models", fake)
    return await md.refresh_provider(db, "gemini")


async def _enabled(db):
    from app.models.ai_catalog import AICatalogModel

    rows = (await db.execute(select(AICatalogModel).where(AICatalogModel.provider_key == "gemini"))).scalars().all()
    return sorted(m.model_key for m in rows if m.enabled)


@pytest.mark.asyncio
async def test_a_list_with_zero_text_models_retires_nothing(db_session, monkeypatch):
    await _gemini(db_session, monkeypatch)
    assert (await _refresh(db_session, monkeypatch, ["gemini-4.0-flash", "gemini-4.0-flash-lite", "gemini-4.0-pro"]))["ok"]
    before = await _enabled(db_session)
    res = await _refresh(db_session, monkeypatch, ["text-embedding-005"])
    assert res["ok"] is False and res["retired"] == 0 and res["guard"] == "no_text_models" and "چیزی خاموش نشد" in res["message"]
    assert await _enabled(db_session) == before


@pytest.mark.asyncio
async def test_retiring_more_than_half_at_once_is_refused_and_reported(db_session, monkeypatch):
    await _gemini(db_session, monkeypatch)
    live = ["gemini-4.0-flash", "gemini-4.0-flash-lite", "gemini-4.0-pro", "gemini-3.8-flash"]
    assert (await _refresh(db_session, monkeypatch, live))["ok"]
    before = await _enabled(db_session)
    res = await _refresh(db_session, monkeypatch, ["gemini-4.0-flash"])  # 3 of 4 would vanish
    assert res["ok"] is False and res["guard"] == "mass_retire" and res["retired"] == 0
    assert sorted(res["would_retire"]) == ["gemini-3.8-flash", "gemini-4.0-flash-lite", "gemini-4.0-pro"]
    assert await _enabled(db_session) == before
    # half or fewer is a normal retirement
    res = await _refresh(db_session, monkeypatch, ["gemini-4.0-flash", "gemini-4.0-flash-lite"])
    assert res["ok"] and res["retired"] == 2


# ---------------------------------------------------------------------------- #9 price before version, per provider
async def _two_providers(db, monkeypatch):
    from app.models.ai_catalog import AICatalogModel, AICatalogProvider

    for k, env in (("gemini", "GEMINI_API_KEY"), ("anthropic", "ANTHROPIC_API_KEY")):
        if await db.get(AICatalogProvider, k) is None:
            from app.services.ai.catalog import seed_ai_catalog

            await seed_ai_catalog(db)
        (await db.get(AICatalogProvider, k)).enabled = True
        monkeypatch.setenv(env, "k-test-1234567890abcdef")
    for m in (await db.execute(select(AICatalogModel))).scalars().all():
        m.enabled = False
    await db.commit()

    def add(key, prov, **kw):
        db.add(AICatalogModel(model_key=key, provider_key=prov, display_name=key, enabled=True, source="custom",
                              capabilities=kw.pop("caps", ["text", "fast"]), priority=kw.pop("priority", 2), **kw))

    return add


@pytest.mark.asyncio
async def test_known_price_breaks_a_tie_before_the_version(db_session, monkeypatch):
    from app.services.ai.manager import ai_manager

    add = await _two_providers(db_session, monkeypatch)
    add("gemini-9.0-flash", "gemini", input_cost_per_1m=0.5, output_cost_per_1m=3.0)   # newer, pricier
    add("gemini-4.0-flash", "gemini", input_cost_per_1m=0.1, output_cost_per_1m=0.4)   # older, cheaper
    await db_session.commit()
    assert (await ai_manager.resolve(db_session, "summarization")).model_key == "gemini-4.0-flash"


@pytest.mark.asyncio
async def test_versions_are_never_compared_across_providers(db_session, monkeypatch):
    from app.services.ai.manager import ai_manager

    add = await _two_providers(db_session, monkeypatch)
    add("gemini-4.0-flash", "gemini")       # first in database order
    add("claude-sonnet-9-9", "anthropic")   # «9.9» is not newer than a Gemini «4.0» — different products
    await db_session.commit()
    assert (await ai_manager.resolve(db_session, "summarization")).model_key == "gemini-4.0-flash"
    # within ONE provider the newer version still wins a tie
    add("gemini-4.5-flash", "gemini")
    await db_session.commit()
    assert (await ai_manager.resolve(db_session, "summarization")).model_key == "gemini-4.5-flash"


# ---------------------------------------------------------------------------- #10 no capable model ⇒ cheaper fast, logged
@pytest.mark.asyncio
async def test_no_capable_model_falls_back_to_the_cheaper_fast_model_and_logs_it(db_session, monkeypatch, caplog):
    from app.services.ai.manager import ai_manager

    add = await _two_providers(db_session, monkeypatch)
    # lla_assistant requires vision+documents; nobody has documents
    add("big-reasoner", "gemini", caps=["text", "vision", "reasoning"], priority=1, input_cost_per_1m=5, output_cost_per_1m=20)
    add("fast-pricey", "gemini", caps=["text", "vision", "fast"], priority=2, input_cost_per_1m=1, output_cost_per_1m=8)
    add("fast-cheap", "anthropic", caps=["text", "vision", "fast"], priority=3, input_cost_per_1m=0.1, output_cost_per_1m=0.4)
    await db_session.commit()
    with caplog.at_level(logging.WARNING, logger="app.services.ai.manager"):
        rm = await ai_manager.resolve(db_session, "lla_assistant")
    assert rm.model_key == "fast-cheap"
    assert any("no enabled model has" in r.getMessage() and "fast-cheap" in r.getMessage() for r in caplog.records)
    order = [m.model_key for m in await ai_manager._auto_order(db_session, "lla_assistant")]
    assert order == ["fast-cheap", "fast-pricey", "big-reasoner"]


@pytest.mark.asyncio
async def test_no_capable_model_without_prices_prefers_the_light_fast_tier(db_session, monkeypatch):
    """OpenAI-only (seed has no prices and no «documents» model): the assistant gets GPT-4o mini, not GPT-4o."""
    from app.models.ai_catalog import AICatalogProvider
    from app.services.ai.catalog import seed_ai_catalog
    from app.services.ai.manager import ai_manager

    await seed_ai_catalog(db_session)
    for k in ("GEMINI_API_KEY", "ANTHROPIC_API_KEY", "CLAUDE_CODE_OAUTH_TOKEN"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-openai-1234567890")
    for p in (await db_session.execute(select(AICatalogProvider))).scalars().all():
        p.enabled = p.key == "openai"
    await db_session.commit()
    assert (await ai_manager.resolve(db_session, "lla_assistant")).model_key == "gpt-4o-mini"
